#!/usr/bin/env python3
"""Reject accidentally staged runtime secrets without echoing their values."""
import io
import json
import re
import subprocess
from pathlib import Path

from dev_env import DATA, ROOT


def secret_values(value, key=""):
    if isinstance(value, dict):
        for name, child in value.items():
            yield from secret_values(child, name)
    elif isinstance(value, list):
        for child in value:
            yield from secret_values(child, key)
    elif isinstance(value, str) and any(word in key.lower() for word in ("password", "token", "api_key", "passphrase")):
        if len(value) >= 8:
            yield value.encode()


# A real identifier is never glued to letters or digits; hashes in lock files are, so they are not matched.
EDGE_BEFORE = r"(?<![A-Za-z0-9])"
EDGE_AFTER = r"(?![A-Za-z0-9])"
TAX_ID = re.compile(EDGE_BEFORE + r"[0-9][0-9A-HJ-NPQRTUWXY]{17}" + EDGE_AFTER)
PHONE = re.compile(EDGE_BEFORE + r"1[3-9][0-9]{9}" + EDGE_AFTER)
EMAIL = re.compile(r"[A-Za-z0-9_.+%-]+@([A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+)")
INVOICE_NUMBER = re.compile(EDGE_BEFORE + r"[0-9]{20}" + EDGE_AFTER)
FICTIONAL_TAX_MARK = "XXXXXXXX"
EMAIL_DOMAINS = ("example", "invalid", "reimb.local")
# Fictional invoice numbers are allowed only where fixtures and tests live.
INVOICE_NUMBER_DIRS = {"fixtures", "tests"}


def document_text(data, path):
    """What a reader of the file sees. Compressed binary bytes would only produce random matches."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True)
        return "\n".join(str(value) for sheet in book for row in sheet.iter_rows(values_only=True) for value in row if value is not None)
    if suffix in (".png", ".jpg", ".jpeg"):
        # Pixels carry no text layer; images are fictional renders reviewed by eye.
        return ""
    return data.decode("utf-8", errors="replace")


def pattern_violations(data, path):
    text = document_text(data, path)
    findings = []
    if any(FICTIONAL_TAX_MARK not in match.group() for match in TAX_ID.finditer(text)):
        findings.append("tax_id")
    if PHONE.search(text):
        findings.append("phone")
    for match in EMAIL.finditer(text):
        domain = match.group(1).lower()
        if not any(domain == allowed or domain.endswith("." + allowed) for allowed in EMAIL_DOMAINS):
            findings.append("email")
    if not INVOICE_NUMBER_DIRS & set(path.parts) and INVOICE_NUMBER.search(text):
        findings.append("invoice_number")
    return sorted(set(findings))


def main():
    local_secrets = []
    for path in (DATA / "credentials.json", DATA / "octos-smoke/profiles/reimb-smoke.json"):
        if path.exists():
            local_secrets.extend(secret_values(json.loads(path.read_text())))
    paths = subprocess.check_output(["git", "ls-files", "-z", "--cached", "--others", "--exclude-standard"], cwd=ROOT).split(b"\0")
    violations = []
    for name in set(paths):
        if not name:
            continue
        path = ROOT / name.decode()
        if not path.is_file():
            continue
        data = path.read_bytes()
        if any(secret in data for secret in local_secrets) or re.search(rb"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----", data):
            violations.append({"file": name.decode(), "rules": ["secret"]})
        rules = pattern_violations(data, Path(name.decode()))
        if rules:
            violations.append({"file": name.decode(), "rules": rules})
    print(json.dumps({"status": "failed" if violations else "passed", "files_with_secret_matches": violations}))
    return bool(violations)


if __name__ == "__main__":
    raise SystemExit(main())
