"""Final review must catch edits to the staged package; packaging must refuse doctored snapshots."""
import copy
import json
import shutil

import pytest
from openpyxl import load_workbook

from conftest import Runtime, canonical_hash, decision, expected
from test_demo import build_demo

DEMO = expected("demo.json")
LEDGER = DEMO["ledger"]["name"]


@pytest.fixture(scope="module")
def baseline(tmp_path_factory):
    runtime = Runtime(tmp_path_factory.mktemp("baseline") / "data")
    return runtime, build_demo(runtime)


@pytest.fixture
def copy_of(baseline, tmp_path):
    """A private copy of the finished demo runtime, so each test can damage it freely."""
    runtime, run = baseline
    shutil.copytree(runtime.root, tmp_path / "data")
    clone = Runtime(tmp_path / "data")
    return clone, run, clone.batch / "staging" / run["digest"]


def verify(runtime, run):
    return runtime.ok("verify", {"snapshot_hash": run["digest"], "manifest_object_id": run["digest"],
                                 "history_snapshot": run["history"]["validated_snapshot"]})


def failed(result):
    return sorted(check["name"] for check in result["checks"] if not check["passed"])


def edit_cell(path, sheet, cell, value):
    book = load_workbook(path)
    book[sheet][cell] = value
    book.save(path)


def test_changed_amount_cell_is_caught(copy_of):
    runtime, run, staging = copy_of
    edit_cell(staging / LEDGER, "明细", "E3", 1600)
    result = verify(runtime, run)
    assert result["passed"] is False and {"row_completeness", "category_totals", "table_disk"} <= set(failed(result))


def test_changed_total_only_is_caught(copy_of):
    runtime, run, staging = copy_of
    edit_cell(staging / LEDGER, "汇总", "B5", 5000)
    assert "category_totals" in failed(verify(runtime, run))


def test_formula_injected_into_workbook_is_caught(copy_of):
    runtime, run, staging = copy_of
    edit_cell(staging / LEDGER, "明细", "E12", "=SUM(E2:E11)")
    result = verify(runtime, run)
    assert result["passed"] is False and "row_completeness" in failed(result)


def test_deleted_attachment_is_caught(copy_of):
    runtime, run, staging = copy_of
    (staging / DEMO["ledger"]["rows"][0]["file"]).unlink()
    assert "table_disk" in failed(verify(runtime, run))


def test_renamed_attachment_amount_is_caught(copy_of):
    runtime, run, staging = copy_of
    original = DEMO["ledger"]["rows"][5]["file"]
    renamed = original.replace("_30.8_", "_300.8_")
    (staging / original).rename(staging / renamed)
    edit_cell(staging / LEDGER, "明细", "I7", renamed)
    result = verify(runtime, run)
    assert {"filename_fields", "row_completeness", "table_disk"} <= set(failed(result))


def test_swapped_attachment_content_is_caught(copy_of):
    runtime, run, staging = copy_of
    first, second = (staging / DEMO["ledger"]["rows"][index]["file"] for index in (3, 4))
    data = first.read_bytes()
    first.write_bytes(second.read_bytes())
    second.write_bytes(data)
    assert "table_disk" in failed(verify(runtime, run))


def test_replaced_attachment_with_rewritten_manifest_is_caught(copy_of):
    """The manifest is not trusted: attachments are checked against the snapshot's source hashes too."""
    import hashlib
    runtime, run, staging = copy_of
    name = DEMO["ledger"]["rows"][3]["file"]
    forged = (staging / DEMO["ledger"]["rows"][4]["file"]).read_bytes()
    (staging / name).write_bytes(forged)
    manifest = json.loads((staging / "manifest.json").read_text())
    for entry in manifest["files"]:
        if entry["relative_name"] == name:
            entry.update(sha256=hashlib.sha256(forged).hexdigest(), bytes=len(forged))
    (staging / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False))
    result = verify(runtime, run)
    assert result["passed"] is False and "table_disk" in failed(result)


def test_consistent_forgery_is_caught_against_the_snapshot(copy_of):
    """Every derived value and the manifest agree with the forged amount; only the snapshot disagrees."""
    import hashlib
    runtime, run, staging = copy_of
    row = next(r for r in DEMO["ledger"]["rows"] if r["item"] == "F01")
    old_name, new_name = row["file"], row["file"].replace("_30.8_", "_130.8_")
    (staging / old_name).rename(staging / new_name)
    book = load_workbook(staging / LEDGER)
    detail, summary = book["明细"], book["汇总"]
    line = next(r for r in range(2, 12) if detail.cell(r, 9).value == old_name)
    detail.cell(line, 5).value = 130.8
    detail.cell(line, 9).value = new_name
    detail.cell(12, 5).value = 5001.4
    summary["B4"], summary["B5"] = 544.3, 5001.4
    book.save(staging / LEDGER)
    manifest = json.loads((staging / "manifest.json").read_text())
    for entry in manifest["rows"]:
        if entry["relative_name"] == old_name:
            entry.update(amount_cents=13080, relative_name=new_name)
            entry["business_hash"] = hashlib.sha256(json.dumps(
                {k: v for k, v in entry.items() if k != "business_hash"}, ensure_ascii=False, sort_keys=True,
                separators=(",", ":")).encode()).hexdigest()
    ledger = (staging / LEDGER).read_bytes()
    for entry in manifest["files"]:
        if entry["relative_name"] == old_name:
            entry["relative_name"] = new_name
        if entry["relative_name"] == LEDGER:
            entry.update(sha256=hashlib.sha256(ledger).hexdigest(), bytes=len(ledger))
    manifest["total_cents"] += 10000
    manifest["category_totals"]["餐饮费"] += 10000
    (staging / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False))
    result = verify(runtime, run)
    assert result["passed"] is False and "Amount differs from snapshot" in result["issues"]


def test_reordered_rows_are_caught(copy_of):
    runtime, run, staging = copy_of
    book = load_workbook(staging / LEDGER)
    detail = book["明细"]
    for column in range(2, 10):
        first, second = detail.cell(2, column).value, detail.cell(3, column).value
        detail.cell(2, column).value, detail.cell(3, column).value = second, first
    book.save(staging / LEDGER)
    assert {"date_order", "row_completeness"} <= set(failed(verify(runtime, run)))


def test_extra_file_is_caught(copy_of):
    runtime, run, staging = copy_of
    (staging / "extra.pdf").write_bytes(b"%PDF-1.4\n")
    assert "table_disk" in failed(verify(runtime, run))


def test_edited_manifest_fails_closed(copy_of):
    runtime, run, staging = copy_of
    manifest = json.loads((staging / "manifest.json").read_text())
    manifest["total_cents"] += 100
    (staging / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False))
    assert verify(runtime, run)["passed"] is False
    manifest["unexpected"] = 1
    (staging / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False))
    assert verify(runtime, run)["passed"] is False


def test_edited_staged_snapshot_fails_closed(copy_of):
    runtime, run, staging = copy_of
    snapshot = json.loads((staging / "snapshot.json").read_text())
    snapshot["applicant"] = "周敏"
    (staging / "snapshot.json").write_text(json.dumps(snapshot, ensure_ascii=False))
    assert verify(runtime, run)["passed"] is False


def test_corrupt_workbook_fails_closed(copy_of):
    runtime, run, staging = copy_of
    (staging / LEDGER).write_bytes(b"PK\x03\x04 not a workbook")
    assert verify(runtime, run)["passed"] is False


def test_history_changed_after_package_fails(copy_of):
    runtime, run, _ = copy_of
    history = copy.deepcopy(run["history"]["validated_snapshot"])[:-1]
    result = runtime.ok("verify", {"snapshot_hash": run["digest"], "manifest_object_id": run["digest"],
                                   "history_snapshot": history})
    assert result["passed"] is False


def repackage(runtime, raw, code, exit_code=3):
    runtime.fails("package", {"confirmed_snapshot": raw, "expected_snapshot_hash": canonical_hash(raw)}, code, exit_code)


def by_item(raw, code):
    return next(entry for entry in raw["items"] if entry["id"] == "item-" + code)


def test_snapshot_hash_must_match(copy_of):
    runtime, run, _ = copy_of
    runtime.fails("package", {"confirmed_snapshot": run["snapshot"], "expected_snapshot_hash": "0" * 64},
                  "SNAPSHOT_MISMATCH", 3)


def test_doctored_extracted_amount_is_refused(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    invoice = by_item(raw, "F01")["invoice"]
    invoice["amount_cents"]["value"] = 13080
    invoice["amount_upper"]["value"] = "壹佰叁拾圆捌角"
    repackage(runtime, raw, "SNAPSHOT_MISMATCH")


def test_candidate_fact_cannot_be_packaged(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    target = by_item(raw, "F09")
    target["invoice"] = run["candidate"]
    repackage(runtime, raw, "FACT_UNCONFIRMED")


def test_confirmation_without_matching_decision_is_refused(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    for entry in raw["decisions"]:
        if entry["kind"] == "confirm_visual":
            entry["source_event_id"] = "event-other"
    repackage(runtime, raw, "FACT_UNCONFIRMED")


def test_replacement_needs_its_decision(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    raw["decisions"] = [entry for entry in raw["decisions"] if entry["kind"] != "replace_unpaid_invoice"]
    by_item(raw, "F10")["decision_ids"] = []
    repackage(runtime, raw, "HISTORY_UNKNOWN")


def test_over_limit_needs_explanation(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    raw["decisions"] = [entry for entry in raw["decisions"] if entry["kind"] != "explain_over_limit"]
    by_item(raw, "F08")["decision_ids"] = []
    repackage(runtime, raw, "EVIDENCE_MISSING")


@pytest.mark.parametrize("mutate", [
    lambda raw: raw["decisions"].append(decision("decision-stray", "item-F01", "reject", {"reason": "x"})),
    lambda raw: raw["decisions"][0].update(item_id="item-F02"),
    lambda raw: raw["decisions"][0].update(expected_revision=99),
    lambda raw: raw["decisions"][0].update(payload={"evidence_id": "x", "action": "approve"}),
    lambda raw: raw["decisions"][0].update(kind="approve"),
    lambda raw: raw["items"].append(copy.deepcopy(raw["items"][0])),
    lambda raw: raw.update(revision=True),
    lambda raw: raw["items"][0]["invoice"]["amount_cents"].update(value=True),
])
def test_malformed_snapshot_is_rejected(copy_of, mutate):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    mutate(raw)
    repackage(runtime, raw, "INVALID_SCHEMA", 2)


def test_duplicate_invoice_in_history_is_refused(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    raw["items"] = [entry for entry in raw["items"] if not entry["decision_ids"]]
    raw["decisions"] = []
    # F12 is already in history; the rejected invoice cannot be smuggled into a snapshot.
    f12 = runtime.ok("extract", {"source_file_id": run["sources"]["F12.pdf"]["id"]})["invoice"]
    raw["items"].append({"id": "item-F12", "invoice": f12, "service_date": "2026-09-28", "category": "打车",
                         "short_name": "滴滴", "expense_detail": "市内交通", "decision_ids": [], "replaces_invoice_no": None})
    repackage(runtime, raw, "DUPLICATE_INVOICE")


def test_wrong_buyer_is_refused(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    raw["items"] = [entry for entry in raw["items"] if not entry["decision_ids"]]
    raw["decisions"] = []
    f11 = runtime.ok("extract", {"source_file_id": run["sources"]["F11.pdf"]["id"]})["invoice"]
    raw["items"].append({"id": "item-F11", "invoice": f11, "service_date": "2026-10-21", "category": "电脑",
                         "short_name": "某某", "expense_detail": "文具", "decision_ids": [], "replaces_invoice_no": None})
    repackage(runtime, raw, "WRONG_BUYER")


def test_policy_change_invalidates_snapshot(copy_of):
    runtime, run, _ = copy_of
    runtime.policy.write_text(runtime.policy.read_text().replace("hotel_per_night_cents: 50000", "hotel_per_night_cents: 60000"))
    runtime.fails("package", {"confirmed_snapshot": run["snapshot"], "expected_snapshot_hash": run["digest"]},
                  "SNAPSHOT_MISMATCH", 3)


def test_restaged_artifact_mismatch_is_refused(copy_of):
    runtime, run, staging = copy_of
    edit_cell(staging / LEDGER, "明细", "D2", "改过")
    runtime.fails("package", {"confirmed_snapshot": run["snapshot"], "expected_snapshot_hash": run["digest"]},
                  "SNAPSHOT_MISMATCH", 3)
