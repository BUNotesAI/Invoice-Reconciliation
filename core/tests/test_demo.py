"""Demo batch through the real CLI: ingest -> extract -> history -> gates -> package -> verify."""
import json

import pytest
from openpyxl import load_workbook

from conftest import (FIXTURES, Runtime, canonical_hash, confirm, confirmed_fact_ids, decision, expected, item,
                      snapshot)

DEMO = expected("demo.json")
CODES = [f"F{number:02d}" for number in range(1, 13)]


def build_demo(runtime):
    """Run the P1 pipeline; service dates and decisions come from the hand-written answers."""
    sources = {}
    for path in sorted((FIXTURES / "demo").iterdir()):
        sources[path.name] = runtime.ingest(runtime.upload(path))
    history = runtime.ok("history", {"action": "validate_import",
                                     "entries": json.loads((FIXTURES / "demo" / "history.json").read_text())})
    invoices, extract_issues = {}, {}
    for code in CODES:
        result = runtime.ok("extract", {"source_file_id": sources[f"{code}.pdf"]["id"]})
        invoices[code], extract_issues[code] = result["invoice"], result["issues"]
    vision = runtime.ok("extract", {"source_file_id": sources["F09.pdf"]["id"],
                                    "vision_candidate": DEMO["vision_candidates"]["F09"]})
    candidate = vision["invoice"]
    listed = [invoices[code] for code in CODES if invoices[code]] + [candidate]
    gates = runtime.ok("gates", {"invoices": listed, "history_snapshot": history["validated_snapshot"]})
    confirmed = confirm(candidate, {}, "event-confirm-f09")
    after = runtime.ok("gates", {"invoices": [confirmed], "history_snapshot": history["validated_snapshot"]})
    invoices["F09"] = confirmed
    decisions = []
    for spec in DEMO["decisions"]:
        payload = spec["payload"]
        if spec["kind"] == "confirm_visual":
            payload = {"fact_ids": confirmed_fact_ids(confirmed)}
        decisions.append(decision(spec["id"], "item-" + spec["item"], spec["kind"], payload,
                                  event="event-confirm-f09" if spec["kind"] == "confirm_visual" else None))
    items = []
    for code, spec in DEMO["package_items"].items():
        ids = [entry["id"] for entry in DEMO["decisions"] if entry["item"] == code]
        items.append(item("item-" + code, invoices[code], spec["service_date"], spec["category"], spec["short_name"],
                          spec["expense_detail"], ids, spec.get("replaces_invoice_no")))
    raw = snapshot(gates["policy_hash"], history["history_hash"], items, decisions, batch_id=DEMO["batch_id"],
                   applicant=DEMO["applicant"], period=DEMO["period"], revision=DEMO["revision"])
    digest = canonical_hash(raw)
    package = runtime.ok("package", {"confirmed_snapshot": raw, "expected_snapshot_hash": digest})
    verify = runtime.ok("verify", {"snapshot_hash": digest, "manifest_object_id": package["manifest_object_id"],
                                   "history_snapshot": history["validated_snapshot"]})
    return dict(sources=sources, history=history, invoices=invoices, extract_issues=extract_issues, candidate=candidate,
                gates=gates, after=after, snapshot=raw, digest=digest, package=package, verify=verify)


@pytest.fixture(scope="module")
def demo(tmp_path_factory):
    runtime = Runtime(tmp_path_factory.mktemp("demo") / "data")
    return runtime, build_demo(runtime)


def test_ingest_detects_every_demo_file(demo):
    _, run = demo
    assert {name: source["detected_type"] for name, source in run["sources"].items()} == DEMO["detected_types"]


def test_extract_matches_hand_written_facts(demo):
    _, run = demo
    for code, want in DEMO["invoices"].items():
        invoice = run["invoices"][code] if code != "F09" else None
        if want is None:
            assert run["extract_issues"][code] == DEMO["extract_issues_without_vision"][code]
            continue
        got = {name: invoice[name]["value"] for name in ("invoice_no", "issue_date", "amount_cents", "amount_upper",
                                                         "buyer_name", "buyer_tax_id", "seller_name")}
        got.update(order_ref=invoice["order_ref"]["value"] if invoice["order_ref"] else None,
                   service_period=invoice["service_period"], issues=run["extract_issues"][code])
        assert got == want, code
        assert all(invoice[name]["level"] == "extracted" and invoice[name]["source"]["locator"]["type"] == "pdf_page"
                   for name in ("invoice_no", "amount_cents", "buyer_tax_id"))


def test_vision_reading_stays_candidate_until_confirmed(demo):
    _, run = demo
    candidate = run["candidate"]
    assert {candidate[name]["level"] for name in ("invoice_no", "amount_cents", "seller_name")} == {"candidate"}
    assert candidate["invoice_no"]["source"]["method"] == "vision"


def test_history_import(demo):
    _, run = demo
    assert len(run["history"]["validated_snapshot"]) == DEMO["history"]["entries"]
    assert run["history"]["issues"] == DEMO["history"]["issues"]


def gate_view(result):
    return {"disposition": result["disposition"], "review_reasons": result["review_reasons"],
            "replacements": [row["invoice_no"] for row in result["replacement_candidates"]]}


def test_gates_match_hand_written_dispositions(demo):
    _, run = demo
    by_number = {row["invoice_no"]: row for row in run["gates"]["items"]}
    for code, want in DEMO["gates"].items():
        number = DEMO["invoices"][code]["invoice_no"] if DEMO["invoices"][code] else DEMO["vision_candidates"][code]["invoice_no"]
        assert gate_view(by_number[number]) == want, code
    assert gate_view(run["after"]["items"][0]) == DEMO["gates_after_confirmation"]["F09"]
    rejected = [row for row in run["gates"]["items"] if row["disposition"] == "rejected"]
    assert len(rejected) == 2


def test_package_rows_totals_and_names(demo):
    runtime, run = demo
    manifest = run["package"]["staging_manifest"]
    ledger = DEMO["ledger"]
    assert manifest["ledger_name"] == ledger["name"]
    assert [(row["item_id"], row["service_date"], row["amount_cents"], row["btype"], row["summary"], row["expense_detail"],
             row["relative_name"]) for row in manifest["rows"]] == [
        ("item-" + row["item"], row["service_date"], row["amount_cents"], row["btype"], row["summary"], row["expense_detail"],
         row["file"]) for row in ledger["rows"]]
    assert manifest["category_totals"] == ledger["category_totals"]
    assert manifest["total_cents"] == ledger["total_cents"] and len(manifest["rows"]) == ledger["entered_count"]
    staging = runtime.batch / "staging" / run["digest"]
    assert sorted(path.name for path in staging.iterdir()) == sorted(
        [row["file"] for row in ledger["rows"]] + [ledger["name"], "manifest.json", "snapshot.json"])


def test_workbook_cells_read_back(demo):
    runtime, run = demo
    book = load_workbook(runtime.batch / "staging" / run["digest"] / DEMO["ledger"]["name"])
    detail, summary = book["明细"], book["汇总"]
    first = [detail.cell(2, column).value for column in range(1, 10)]
    row = DEMO["ledger"]["rows"][0]
    assert first == [1, row["btype"], row["summary"], row["expense_detail"], 1460, row["service_date"],
                     "26112000000300012801", "中国国际航空股份有限公司", row["file"]]
    assert detail.cell(12, 1).value == "合计" and detail.cell(12, 5).value == 4901.4
    assert [[summary.cell(r, c).value for c in (1, 2)] for r in range(1, 6)] == [
        ["报销类型", "金额"], ["差旅-交通费", 2897.1], ["差旅-住宿费", 1560], ["餐饮费", 444.3], ["合计", 4901.4]]


def test_verify_passes_all_six_checks(demo):
    _, run = demo
    assert run["verify"]["passed"] is DEMO["verify"]["passed"]
    assert [check["name"] for check in run["verify"]["checks"] if check["passed"]] == DEMO["verify"]["checks"]


def test_package_is_idempotent_and_byte_stable(demo, tmp_path):
    runtime, run = demo
    again = runtime.ok("package", {"confirmed_snapshot": run["snapshot"], "expected_snapshot_hash": run["digest"]})
    assert again == run["package"]
    other = Runtime(tmp_path / "data")
    second = build_demo(other)
    assert second["digest"] == run["digest"]
    first_files = {entry["relative_name"]: entry["sha256"] for entry in run["package"]["staging_manifest"]["files"]}
    second_files = {entry["relative_name"]: entry["sha256"] for entry in second["package"]["staging_manifest"]["files"]}
    assert first_files == second_files
    ledger = DEMO["ledger"]["name"]
    assert (runtime.batch / "staging" / run["digest"] / ledger).read_bytes() == \
        (other.batch / "staging" / second["digest"] / ledger).read_bytes()
