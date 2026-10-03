"""One counterexample per guard found weak in the P1 adversarial review (A1-A9, M01-M20).

Each test asserts the guard's own code or issue text, so a guard cannot pass because another check fired first."""
import copy
import json
import os
import shutil

import pytest
from openpyxl import load_workbook

from conftest import (FIXTURES, Runtime, canonical_hash, confirm, confirmed_fact_ids, decision, expected,
                      item as make_item)
from test_demo import build_demo
from test_hardening import minimal_pdf

DEMO = expected("demo.json")
LEDGER = DEMO["ledger"]["name"]


@pytest.fixture(scope="module")
def baseline(tmp_path_factory):
    runtime = Runtime(tmp_path_factory.mktemp("guards") / "data")
    return runtime, build_demo(runtime)


@pytest.fixture
def copy_of(baseline, tmp_path):
    runtime, run = baseline
    shutil.copytree(runtime.root, tmp_path / "data")
    clone = Runtime(tmp_path / "data")
    return clone, run, clone.batch / "staging" / run["digest"]


def by_item(raw, code):
    return next(entry for entry in raw["items"] if entry["id"] == "item-" + code)


def package_fails(runtime, raw, code, exit_code=3):
    runtime.fails("package", {"confirmed_snapshot": raw, "expected_snapshot_hash": canonical_hash(raw)}, code, exit_code)


def verify(runtime, digest, history):
    return runtime.ok("verify", {"snapshot_hash": digest, "manifest_object_id": digest, "history_snapshot": history})


def issues(result):
    return result["issues"]


def extracted(runtime, name):
    source = runtime.ingest(runtime.upload(FIXTURES / "edge" / name))
    return runtime.ok("extract", {"source_file_id": source["id"]})["invoice"]


def staged_by_faulty_packer(runtime, raw):
    """Simulates a packaging defect: the real renderer writes staging with every snapshot guard removed."""
    from reimb_core import package as packer
    from reimb_core.models import Snapshot
    from reimb_core.policy import load_policy
    from reimb_core.storage import Store
    previous = os.environ.get("REIMB_DATA")
    os.environ["REIMB_DATA"] = str(runtime.root)
    original = packer.validate_snapshot

    def lenient(raw_, _expected, _policy, store_):
        snapshot = Snapshot.model_validate(raw_)
        return snapshot, {entry.id: store_.source(entry.invoice.source_file_id) for entry in snapshot.items}
    packer.validate_snapshot = lenient
    try:
        digest = canonical_hash(raw)
        packer.package(Store(str(runtime.batch), runtime.root), load_policy(runtime.policy), raw, digest)
    finally:
        packer.validate_snapshot = original
        if previous is None:
            os.environ.pop("REIMB_DATA", None)
        else:
            os.environ["REIMB_DATA"] = previous
    return digest


# --- A1: lodging rules follow the facts ---------------------------------------------------------------------


def test_hotel_invoice_relabelled_as_flight_is_refused(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    raw["decisions"] = [d for d in raw["decisions"] if d["kind"] != "explain_over_limit"]
    by_item(raw, "F08").update(category="机票", decision_ids=[])
    package_fails(runtime, raw, "FIELD_CONFLICT")


def test_flight_labelled_as_lodging_is_refused(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    by_item(raw, "F07")["category"] = "住宿"
    package_fails(runtime, raw, "FIELD_CONFLICT")


def test_stay_service_date_must_be_check_in(copy_of):  # M10
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    by_item(raw, "F08")["service_date"] = "2026-10-07"
    package_fails(runtime, raw, "FIELD_CONFLICT")


# --- A2: stay dates are facts ---------------------------------------------------------------------------------


@pytest.fixture
def hotel_image(copy_of):
    runtime, run, _ = copy_of
    source = runtime.ingest(runtime.upload(FIXTURES / "edge" / "E18.pdf"))
    reading = {"invoice_no": "26112000000400099991", "issue_date": "2026-10-09", "amount_cents": "1560.00",
               "amount_upper": "壹仟伍佰陆拾圆整", "buyer_name": "示例科技有限公司", "buyer_tax_id": "91440300XXXXXXXX0A",
               "seller_name": "北京燕园会展酒店有限公司", "project": "*住宿服务*住宿费", "remark": "入住 2026-10-06 离店 2026-10-09 3晚"}
    candidate = runtime.ok("extract", {"source_file_id": source["id"], "vision_candidate": reading})
    return runtime, run, candidate


def with_hotel(run, invoice, decision_ids, decisions, service_date="2026-10-06"):
    raw = copy.deepcopy(run["snapshot"])
    raw["decisions"] += decisions
    raw["items"].append(make_item("item-HOTEL", invoice, service_date, "住宿", "燕园会展酒店", "住宿", decision_ids))
    return raw


def test_image_stay_is_a_candidate(hotel_image):
    _, _, candidate = hotel_image
    period = candidate["invoice"]["service_period"]
    assert candidate["issues"] == ["FACT_UNCONFIRMED"] and {fact["level"] for fact in period.values()} == {"candidate"}


def test_forged_plain_stay_period_is_refused(hotel_image):
    runtime, run, candidate = hotel_image
    confirmed = confirm(candidate["invoice"], {}, "event-hotel")
    confirmed["service_period"] = {"check_in": "2026-10-05", "check_out": "2026-10-09", "nights": 4}
    raw = with_hotel(run, confirmed, [], [])
    package_fails(runtime, raw, "INVALID_SCHEMA", 2)


def test_stay_facts_outside_the_confirmation_are_refused(hotel_image):
    runtime, run, candidate = hotel_image
    confirmed = confirm(candidate["invoice"], {}, "event-hotel")
    ids = [fact_id for fact_id in confirmed_fact_ids(confirmed) if ".nights." not in fact_id]
    explain = decision("d-hotel-explain", "item-HOTEL", "explain_over_limit", {"explanation": "会展"})
    confirmation = decision("d-hotel", "item-HOTEL", "confirm_visual", {"fact_ids": ids}, event="event-hotel")
    raw = with_hotel(run, confirmed, ["d-hotel", "d-hotel-explain"], [confirmation, explain])
    package_fails(runtime, raw, "FACT_UNCONFIRMED")


def test_confirmed_image_stay_over_limit_needs_explanation(hotel_image):
    runtime, run, candidate = hotel_image
    confirmed = confirm(candidate["invoice"], {}, "event-hotel")
    confirmation = decision("d-hotel", "item-HOTEL", "confirm_visual", {"fact_ids": confirmed_fact_ids(confirmed)},
                            event="event-hotel")
    package_fails(runtime, with_hotel(run, confirmed, ["d-hotel"], [confirmation]), "EVIDENCE_MISSING")
    explain = decision("d-hotel-explain", "item-HOTEL", "explain_over_limit", {"explanation": "会展"})
    raw = with_hotel(run, confirmed, ["d-hotel", "d-hotel-explain"], [confirmation, explain])
    digest = canonical_hash(raw)
    runtime.ok("package", {"confirmed_snapshot": raw, "expected_snapshot_hash": digest})
    assert verify(runtime, digest, run["history"]["validated_snapshot"])["passed"] is True


# --- A3: verify recomputes the snapshot against the originals --------------------------------------------------


def lenient_verify(runtime, run, raw):
    digest = staged_by_faulty_packer(runtime, raw)
    return verify(runtime, digest, run["history"]["validated_snapshot"])


def test_verify_catches_doctored_amount_without_package(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    f01 = by_item(raw, "F01")["invoice"]
    f01["amount_cents"]["value"], f01["amount_upper"]["value"] = 13080, "壹佰叁拾圆捌角"
    assert "item-F01: original differs (amount)" in issues(lenient_verify(runtime, run, raw))


def test_verify_catches_wrong_buyer_without_package(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    f11 = runtime.ok("extract", {"source_file_id": run["sources"]["F11.pdf"]["id"]})["invoice"]
    raw["items"].append(make_item("item-F11", f11, "2026-10-21", "电脑", "某某", "文具"))
    assert "item-F11: buyer differs from policy" in issues(lenient_verify(runtime, run, raw))


def test_verify_catches_invoice_already_in_history(copy_of):  # M01
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    f12 = runtime.ok("extract", {"source_file_id": run["sources"]["F12.pdf"]["id"]})["invoice"]
    raw["items"].append(make_item("item-F12", f12, "2026-09-28", "打车", "滴滴", "市内交通"))
    assert "Invoice number already in history" in issues(lenient_verify(runtime, run, raw))


def test_verify_catches_duplicate_number_in_batch(copy_of):  # M20
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    # A different short name keeps the attachment names apart, so only the number repeats.
    raw["items"].append(make_item("item-E17", extracted(runtime, "E17.pdf"), "2026-10-13", "餐饮", "瑞幸咖啡", "工作餐"))
    assert "Duplicate invoice number in batch" in issues(lenient_verify(runtime, run, raw))


def test_verify_catches_unexplained_stay_without_package(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    raw["decisions"] = [d for d in raw["decisions"] if d["kind"] != "explain_over_limit"]
    by_item(raw, "F08")["decision_ids"] = []
    assert "item-F08: over-limit stay without explanation" in issues(lenient_verify(runtime, run, raw))


def test_verify_catches_broken_replacement_without_package(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    raw["decisions"] = [d for d in raw["decisions"] if d["kind"] != "replace_unpaid_invoice"]
    by_item(raw, "F10")["decision_ids"] = []
    assert "item-F10: replacement breaks the history rules" in issues(lenient_verify(runtime, run, raw))


def test_verify_catches_category_against_facts_without_package(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    by_item(raw, "F07")["category"] = "打车"
    assert "item-F07: category differs from invoice facts" in issues(lenient_verify(runtime, run, raw))


# --- A4: what finance sees is bound to the snapshot -----------------------------------------------------------


def rewrite_manifest_hash(staging):
    import hashlib
    manifest = json.loads((staging / "manifest.json").read_text())
    data = (staging / LEDGER).read_bytes()
    for entry in manifest["files"]:
        if entry["relative_name"] == LEDGER:
            entry.update(sha256=hashlib.sha256(data).hexdigest(), bytes=len(data))
    (staging / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False))


@pytest.mark.parametrize("tamper", [
    lambda book: setattr(book["明细"]["E3"], "number_format", '"1,600.00"'),
    lambda book: setattr(book["明细"].row_dimensions[10], "hidden", True),
    lambda book: setattr(book["明细"].column_dimensions["D"], "hidden", True),
    lambda book: book["明细"].__setitem__("I12", "另附现金收据，合计以此为准"),
    lambda book: book["汇总"].__setitem__("C3", "含超标"),
])
def test_display_tampering_is_caught(copy_of, tamper):
    runtime, run, staging = copy_of
    book = load_workbook(staging / LEDGER)
    tamper(book)
    book.save(staging / LEDGER)
    rewrite_manifest_hash(staging)
    result = verify(runtime, run["digest"], run["history"]["validated_snapshot"])
    assert "Workbook differs from the snapshot rendering" in issues(result)


@pytest.mark.parametrize("tamper, issue", [
    (lambda book: book["明细"].__setitem__("G2", 26112000000300012801), "Untrusted text is not a text cell"),  # M02
    (lambda book: setattr(book["明细"]["G2"], "number_format", "General"), "Invoice number format changed"),  # M03
    (lambda book: book["明细"].merge_cells("F12:G12"), "Unexpected merged cells"),  # M04
    (lambda book: book["汇总"].__setitem__("D2", "周敏"), "Applicant differs"),  # M05
    (lambda book: book["汇总"].__setitem__("A2", "差旅费"), "Category total differs"),  # M16
])
def test_each_cell_guard_names_its_issue(copy_of, tamper, issue):
    runtime, run, staging = copy_of
    book = load_workbook(staging / LEDGER)
    tamper(book)
    book.save(staging / LEDGER)
    assert issue in issues(verify(runtime, run["digest"], run["history"]["validated_snapshot"]))


def test_short_name_in_file_name_is_checked(copy_of):  # M06
    runtime, run, staging = copy_of
    old = next(row["file"] for row in DEMO["ledger"]["rows"] if row["item"] == "F01")
    new = old.replace("_瑞幸_", "_瑞星_")
    (staging / old).rename(staging / new)
    book = load_workbook(staging / LEDGER)
    detail = book["明细"]
    line = next(r for r in range(2, 12) if detail.cell(r, 9).value == old)
    detail.cell(line, 9).value = new
    book.save(staging / LEDGER)
    assert "Filename differs from cells" in issues(verify(runtime, run["digest"], run["history"]["validated_snapshot"]))


# --- Re-issue and duplicate guards in package and gates (M07-M09, M11, M12, M15, M18, M19) -------------------


def test_replacement_date_must_equal_original(copy_of):  # M07
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    by_item(raw, "F10")["service_date"] = "2026-09-23"
    package_fails(runtime, raw, "HISTORY_UNKNOWN")


def test_replacement_amount_must_equal_original(copy_of):  # M08
    runtime, run, _ = copy_of
    entries = json.loads((FIXTURES / "demo" / "history.json").read_text())
    for entry in entries:
        if entry["invoice_no"] == "26112000000300002208":
            entry["amount_cents"] = 150000
    history = runtime.ok("history", {"action": "validate_import", "entries": entries})
    raw = copy.deepcopy(run["snapshot"])
    raw["history_hash"] = history["history_hash"]
    package_fails(runtime, raw, "HISTORY_UNKNOWN")


def test_one_original_is_replaced_once(copy_of):  # M09
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    second = decision("d-e16", "item-E16", "replace_unpaid_invoice", {"invoice_no": "26112000000300002208"})
    raw["decisions"].append(second)
    raw["items"].append(make_item("item-E16", extracted(runtime, "E16.pdf"), "2026-09-22", "机票", "国航", "出差",
                                  ["d-e16"], "26112000000300002208"))
    package_fails(runtime, raw, "LINK_CONFLICT")


def test_confirmation_actor_must_be_the_confirming_user(copy_of):  # M11
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    for entry in raw["decisions"]:
        if entry["kind"] == "confirm_visual":
            entry["actor"] = "@someone-else:reimb.local"
    package_fails(runtime, raw, "FACT_UNCONFIRMED")


def test_confirmation_locator_must_match_provenance(copy_of):  # M12
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    by_item(raw, "F09")["invoice"]["invoice_no"]["source"]["locator"]["event_id"] = "event-other"
    package_fails(runtime, raw, "INVALID_SCHEMA", 2)


def test_one_attachment_serves_one_item(copy_of):  # M15
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    f09 = by_item(raw, "F09")
    twin = confirm(run["candidate"], {"invoice_no": "26442000000500010192"}, "event-twin")
    twin_decision = decision("d-twin", "item-TWIN", "confirm_visual", {"fact_ids": confirmed_fact_ids(twin)}, event="event-twin")
    raw["decisions"].append(twin_decision)
    raw["items"].append(make_item("item-TWIN", twin, f09["service_date"], "餐饮", "潮海居", "接待", ["d-twin"]))
    package_fails(runtime, raw, "LINK_CONFLICT")


def test_two_copies_of_one_invoice_are_both_rejected_at_gates(copy_of):  # M18
    runtime, run, _ = copy_of
    f01 = runtime.ok("extract", {"source_file_id": run["sources"]["F01.pdf"]["id"]})["invoice"]
    result = runtime.ok("gates", {"invoices": [f01, extracted(runtime, "E17.pdf")],
                                  "history_snapshot": run["history"]["validated_snapshot"]})
    assert [(row["disposition"], row["review_reasons"]) for row in result["items"]] == [("rejected", ["DUPLICATE_INVOICE"])] * 2


def test_two_copies_of_one_invoice_cannot_be_packaged(copy_of):  # M19
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    raw["items"].append(make_item("item-E17", extracted(runtime, "E17.pdf"), "2026-10-13", "餐饮", "瑞幸", "工作餐"))
    package_fails(runtime, raw, "DUPLICATE_INVOICE")


# --- History knowledge and ordering (M13, M17, A6) ------------------------------------------------------------


def test_imported_initial_payment_counts_as_paid(runtime):  # M13
    entries = json.loads((FIXTURES / "edge" / "history.json").read_text())
    for entry in entries:
        if entry["invoice_no"] == "26112000000700008802":
            entry["initial_paid"] = True
    history = runtime.ok("history", {"action": "validate_import", "entries": entries})
    source = runtime.ingest(runtime.upload(FIXTURES / "edge" / "E02.pdf"))
    invoice = runtime.ok("extract", {"source_file_id": source["id"]})["invoice"]
    gates = runtime.ok("gates", {"invoices": [invoice], "history_snapshot": history["validated_snapshot"]})
    assert gates["items"][0]["review_reasons"] == ["HISTORY_ALREADY_PAID"]


def history_with_times(times):
    entry = json.loads((FIXTURES / "demo" / "history.json").read_text())[0]
    entry["events"] = [dict(event, at=at) for event, at in zip(entry["events"], times)]
    return entry


def test_events_out_of_order_are_refused(runtime):  # M17
    entry = history_with_times(["2026-10-03T00:00:00Z", "2026-10-02T00:00:00Z", "2026-10-04T00:00:00Z"])
    runtime.fails("history", {"action": "validate_import", "entries": [entry]}, "INVALID_SCHEMA", 2)


def test_fractional_seconds_are_ordered_by_time(runtime):  # A6
    in_order = history_with_times(["2026-10-01T00:00:00Z", "2026-10-01T00:00:00.5Z", "2026-10-01T00:00:00.9Z"])
    runtime.ok("history", {"action": "validate_import", "entries": [in_order]})
    out_of_order = history_with_times(["2026-10-01T00:00:00.5Z", "2026-10-01T00:00:00Z", "2026-10-01T00:00:01Z"])
    runtime.fails("history", {"action": "validate_import", "entries": [out_of_order]}, "INVALID_SCHEMA", 2)


# --- Upload and PDF hardening (A7, A8) and label names (A9) --------------------------------------------------


def test_hard_link_from_outside_is_refused(runtime, tmp_path):  # A7
    outside = tmp_path / "outside.txt"
    outside.write_text("outside the data root\n")
    link = runtime.uploads / "hl.txt"
    os.link(outside, link)
    runtime.fails("ingest", {"source_path": str(link), "original_name": "hl.txt"}, "INVALID_PATH", 2)


def form_pdf(inner):
    """A page whose only image sits inside a Form XObject, optionally two forms deep."""
    image = b"<< /Type /XObject /Subtype /Image /Width 8000 /Height 5001 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Length 0 >>\nstream\n\nendstream"
    form = b"<< /Type /XObject /Subtype /Form /BBox [0 0 1 1] /Resources << /XObject << /Im1 6 0 R >> >> /Length 0 >>\nstream\n\nendstream"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
               b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /XObject << /Fm1 5 0 R >> >> /Contents 4 0 R >>",
               b"<< /Length %d >>\nstream\n" % len(b"q /Fm1 Do Q") + b"q /Fm1 Do Q" + b"\nendstream", form, image]
    if inner:
        objects[4] = form.replace(b"/Im1 6 0 R", b"/Fm2 7 0 R")
        objects.append(form)
    output, offsets = bytearray(b"%PDF-1.4\n"), []
    for number, body in enumerate(objects, 1):
        offsets.append(len(output))
        output += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(output)
    output += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    output += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    output += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref)
    return bytes(output)


@pytest.mark.parametrize("inner", [False, True])
def test_image_inside_form_xobject_is_bounded(runtime, inner):  # A8
    path = runtime.uploads / "form.pdf"
    path.write_bytes(form_pdf(inner))
    runtime.fails("ingest", {"source_path": str(path), "original_name": "form.pdf"}, "PARSE_LIMIT", 2)


def test_inline_image_is_bounded(runtime):  # A8
    path = runtime.uploads / "inline.pdf"
    path.write_bytes(minimal_pdf(b"q BI /W 8000 /H 5001 /CS /RGB /BPC 8 ID \x00\x00\x00 EI Q"))
    runtime.fails("ingest", {"source_path": str(path), "original_name": "inline.pdf"}, "PARSE_LIMIT", 2)


@pytest.mark.parametrize("short", ["瑞幸_130.8", "瑞幸8", "8瑞幸", "a/b", "瑞 幸"])
def test_label_names_cannot_fake_file_name_fields(copy_of, short):  # A9
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    by_item(raw, "F01")["short_name"] = short
    package_fails(runtime, raw, "INVALID_PATH", 2)


# --- Snapshot hash is taken over the normalised snapshot ---------------------------------------------------


def test_omitted_defaults_name_the_same_snapshot(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    for entry in raw["items"]:
        if entry["replaces_invoice_no"] is None:
            del entry["replaces_invoice_no"]
    # The raw text differs, but the normalised snapshot and its execution key do not.
    assert canonical_hash(raw) != run["digest"]
    again = runtime.ok("package", {"confirmed_snapshot": raw, "expected_snapshot_hash": run["digest"]})
    assert again == run["package"]
    runtime.fails("package", {"confirmed_snapshot": raw, "expected_snapshot_hash": canonical_hash(raw)}, "SNAPSHOT_MISMATCH", 3)



# --- Guards isolated after the first full mutation run -----------------------------------------------------


def test_directory_symlink_inside_the_root_is_refused(runtime):
    real = runtime.upload(FIXTURES / "demo" / "F01.pdf")
    (runtime.root / "alias").symlink_to(runtime.uploads)
    runtime.fails("ingest", {"source_path": str(runtime.root / "alias" / real.name), "original_name": "x"}, "INVALID_PATH", 2)


def test_manifest_business_fields_are_checked(copy_of):
    import hashlib
    runtime, run, staging = copy_of
    manifest = json.loads((staging / "manifest.json").read_text())
    row = manifest["rows"][0]
    row["expense_detail"] = "改过的明细"
    body = {key: value for key, value in row.items() if key != "business_hash"}
    row["business_hash"] = hashlib.sha256(json.dumps(body, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()).hexdigest()
    (staging / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False))
    assert "Manifest business fields differ" in issues(verify(runtime, run["digest"], run["history"]["validated_snapshot"]))


def test_verify_catches_attachment_used_twice(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    twin = confirm(run["candidate"], {"invoice_no": "26442000000500010192"}, "event-twin")
    raw["decisions"].append(decision("d-twin", "item-TWIN", "confirm_visual", {"fact_ids": confirmed_fact_ids(twin)}, event="event-twin"))
    raw["items"].append(make_item("item-TWIN", twin, "2026-10-19", "餐饮", "潮海居", "接待", ["d-twin"]))
    assert "item-TWIN: attachment used twice" in issues(lenient_verify(runtime, run, raw))


def test_verify_catches_unconfirmed_image_reading(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    by_item(raw, "F09")["invoice"] = run["candidate"]
    assert "item-F09: visual reading not confirmed" in issues(lenient_verify(runtime, run, raw))


def test_verify_catches_text_facts_that_are_not_extracted(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    fact = by_item(raw, "F01")["invoice"]["remark"]
    fact.update(level="candidate", source=dict(fact["source"], method="vision",
                                               locator={"type": "image_region", "field": "remark", "region": None}))
    assert "item-F01: text facts not extracted" in issues(lenient_verify(runtime, run, raw))


def test_verify_catches_an_original_replaced_twice(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    raw["decisions"].append(decision("d-e16", "item-E16", "replace_unpaid_invoice", {"invoice_no": "26112000000300002208"}))
    raw["items"].append(make_item("item-E16", extracted(runtime, "E16.pdf"), "2026-09-22", "机票", "国航", "出差",
                                  ["d-e16"], "26112000000300002208"))
    assert "item-E16: original replaced twice" in issues(lenient_verify(runtime, run, raw))


def test_verify_catches_a_reissue_entered_as_new(copy_of):
    runtime, run, _ = copy_of
    raw = copy.deepcopy(run["snapshot"])
    raw["decisions"] = [d for d in raw["decisions"] if d["kind"] != "replace_unpaid_invoice"]
    by_item(raw, "F10").update(decision_ids=[], replaces_invoice_no=None)
    assert "item-F10: re-issue of a history expense without replacement" in issues(lenient_verify(runtime, run, raw))
