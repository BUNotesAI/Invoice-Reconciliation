"""Independent verifier: reread the workbook, the attachments and the originals, never cached totals.

It shares only pure helpers with packaging (value parsing, the workbook renderer, fact-kind rules); every check
recomputes from the frozen snapshot and the stored originals, so a packaging defect cannot pass silently."""
import hashlib
import re
import zipfile
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from .errors import CoreError, require
from .history import validated_history
from .ledger import workbook
from .models import HistoryEntry, Manifest, Snapshot
from .rules import check_category, stay_over_limit
from .storage import inside, pdf_pages
from .values import amount_text, decimal_cents, digest, strict_json, sum_cents

CHECKS = ("row_completeness", "date_order", "unique_invoices", "table_disk", "filename_fields", "category_totals",
          "source_facts")
TWENTY_DIGITS = re.compile(r"(?<![0-9])[0-9]{20}(?![0-9])")
LOWER_AMOUNT = re.compile(r"价税合计小写\s*[:：]\s*[¥￥]?\s*([0-9]+(?:\.[0-9]{1,2})?)")
BUYER_NAME = re.compile(r"购买方名称\s*[:：][ \t]*([^\n]*)")
BUYER_TAX = re.compile(r"购买方税号\s*[:：][ \t]*([^\n]*)")
STAY = re.compile(r"入住\s*([0-9]{4}-[0-9]{2}-[0-9]{2})\s*离店\s*([0-9]{4}-[0-9]{2}-[0-9]{2})")
HEADERS = ["序号", "报销类型", "项目", "费用明细", "金额", "日期", "发票号", "公司全称", "备注"]
RECORDS = {"manifest.json", "snapshot.json"}


class Failure(Exception):
    pass


def actual_cents(cell):
    if cell.data_type != "n" or type(cell.value) not in (int, float):
        raise Failure("Amount cell must be numeric")
    try:
        amount = Decimal(str(cell.value)) * 100
    except InvalidOperation:
        raise Failure("Invalid amount cell") from None
    if not amount.is_finite() or amount != amount.to_integral_value() or amount < 0:
        raise Failure("Amount cell has fractional cents")
    return int(amount)


def neutral(value):
    value = str(value)
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def text_layer_agrees(invoice, text):
    """A separate reading of the original: the number, lower amount, buyer and stay must match the snapshot."""
    problems = []
    if invoice.invoice_no.value not in set(TWENTY_DIGITS.findall(text)):
        problems.append("invoice number")
    amounts = LOWER_AMOUNT.findall(text)
    if len(amounts) != 1 or decimal_cents(amounts[0]) != invoice.amount_cents.value:
        problems.append("amount")
    names, taxes = BUYER_NAME.findall(text), BUYER_TAX.findall(text)
    if [name.strip() for name in names] != [invoice.buyer_name.value] or [tax.strip() for tax in taxes] != [invoice.buyer_tax_id.value]:
        problems.append("buyer")
    stay = STAY.search(text)
    period = invoice.service_period
    if (stay is None) != (period is None) or (stay and (stay[1], stay[2]) != (period.check_in.value, period.check_out.value)):
        problems.append("stay")
    return problems


def source_checks(store, policy, snapshot, history, check):
    """Snapshot against the originals and the policy, recomputed here rather than trusted from packaging."""
    decisions = {decision.id: decision for decision in snapshot.decisions}
    by_number = {entry.invoice_no: entry for entry in history}
    attachments, replaced = set(), set()
    for item in snapshot.items:
        invoice, label = item.invoice, item.id
        applicable = [decisions[key] for key in item.decision_ids]
        check(invoice.source_file_id not in attachments, "source_facts", f"{label}: attachment used twice")
        attachments.add(invoice.source_file_id)
        source, path = store.source(invoice.source_file_id)
        facts = invoice.facts()
        if source.detected_type == "invoice_pdf":
            check(all(fact.level == "extracted" for _, fact in facts), "source_facts", f"{label}: text facts not extracted")
            problems = text_layer_agrees(invoice, "\n".join(pdf_pages(path.read_bytes())))
            check(not problems, "source_facts", f"{label}: original differs ({', '.join(problems)})")
        elif source.detected_type == "image_invoice_pdf":
            confirmations = [d for d in applicable if d.kind == "confirm_visual"]
            check(all(fact.level == "confirmed" and any(
                d.source_event_id == fact.confirmation_event and d.actor == fact.confirmed_by and fact.id in d.typed().fact_ids
                for d in confirmations) for _, fact in facts), "source_facts", f"{label}: visual reading not confirmed")
        else:
            check(False, "source_facts", f"{label}: attachment is not an invoice PDF")
        check(invoice.buyer_name.value == policy.company.name and invoice.buyer_tax_id.value == policy.company.tax_id,
              "source_facts", f"{label}: buyer differs from policy")
        category = policy.categories.get(item.category)
        try:
            shown = check_category(invoice, category) if category else None
        except CoreError:
            shown = None
            check(False, "source_facts", f"{label}: category differs from invoice facts")
        check(category is not None, "source_facts", f"{label}: unknown category")
        kinds = {d.kind for d in applicable}
        if shown == "lodging":
            period = invoice.service_period
            check(period is not None and item.service_date == period.check_in.value, "source_facts",
                  f"{label}: stay dates missing or service date is not check-in")
            check(not stay_over_limit(invoice, policy) or "explain_over_limit" in kinds, "source_facts",
                  f"{label}: over-limit stay without explanation")
        order = invoice.order_ref.value if invoice.order_ref is not None else None
        related = [entry for entry in history if order and entry.order_ref == order and entry.invoice_no != invoice.invoice_no.value]
        original = by_number.get(item.replaces_invoice_no) if item.replaces_invoice_no else None
        if item.replaces_invoice_no is None:
            check(not related, "source_facts", f"{label}: re-issue of a history expense without replacement")
        else:
            check(original is not None and original in related and original.current_status == "voided"
                  and original.payment_knowledge() == "unpaid" and original.amount_cents == invoice.amount_cents.value
                  and original.service_date == item.service_date and len(related) == 1
                  and any(d.kind == "replace_unpaid_invoice" and d.typed().invoice_no == item.replaces_invoice_no for d in applicable),
                  "source_facts", f"{label}: replacement breaks the history rules")
            check(item.replaces_invoice_no not in replaced, "source_facts", f"{label}: original replaced twice")
            replaced.add(item.replaces_invoice_no)


def expected_rows(snapshot, policy):
    """Ledger rows derived here from the snapshot, in service-date order."""
    rows = []
    for item in sorted(snapshot.items, key=lambda entry: (entry.service_date, entry.id)):
        invoice, rule = item.invoice, policy.categories[item.category]
        detail = item.expense_detail
        if item.replaces_invoice_no:
            detail += f"；重开票，原票 {item.replaces_invoice_no} 已作废"
        name = f"{invoice.invoice_no.value}_{item.short_name}_{amount_text(invoice.amount_cents.value)}_{item.category}.pdf"
        rows.append(dict(btype=rule.btype, summary=rule.summary, expense_detail=detail, amount_cents=invoice.amount_cents.value,
                         service_date=item.service_date, invoice_no=invoice.invoice_no.value,
                         seller_name=invoice.seller_name.value, relative_name=name))
    return rows


def verify(store, policy, snapshot_hash, manifest_object_id, history_snapshot):
    require(isinstance(snapshot_hash, str) and re.fullmatch(r"[0-9a-f]{64}", snapshot_hash)
            and manifest_object_id == snapshot_hash, "SNAPSHOT_MISMATCH", "Invalid manifest binding", 3)
    directory = inside(store.batch / "staging" / manifest_object_id, store.root)
    failures = {key: [] for key in CHECKS}

    def check(condition, group, issue):
        if not condition:
            failures[group].append(issue)

    try:
        raw_snapshot = strict_json(inside(directory / "snapshot.json", store.root).read_bytes())
        if digest(raw_snapshot) != snapshot_hash:
            raise Failure("Staged snapshot changed")
        snapshot = Snapshot.model_validate(raw_snapshot)
        manifest = Manifest.model_validate(strict_json(inside(directory / "manifest.json", store.root).read_bytes()))
        if not (manifest.snapshot_hash == snapshot_hash and manifest.policy_hash == policy.sha() == snapshot.policy_hash
                and manifest.history_hash == snapshot.history_hash and manifest.batch_id == snapshot.batch_id
                and manifest.revision == snapshot.revision):
            raise Failure("Manifest binding changed")
        history = validated_history(history_snapshot)
        if history["history_hash"] != snapshot.history_hash:
            raise Failure("History binding changed")
        history_numbers = {entry["invoice_no"] for entry in history["validated_snapshot"]}
        source_checks(store, policy, snapshot, [HistoryEntry.model_validate(row) for row in history["validated_snapshot"]], check)

        ordered = sorted(snapshot.items, key=lambda item: (item.service_date, item.id))
        expected_total = sum_cents(item.invoice.amount_cents.value for item in ordered)
        ledger = policy.naming.ledger.replace("{applicant}", snapshot.applicant).replace("{total}", amount_text(expected_total))
        check(manifest.ledger_name == ledger, "table_disk", "Ledger name differs from policy")
        path = inside(directory / manifest.ledger_name, store.root)
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) >= 200 or sum(entry.file_size for entry in entries) > 32 * 1024 * 1024:
                raise Failure("Workbook archive too large")
        book = load_workbook(path, data_only=False, keep_links=False)
        try:
            if book.sheetnames != ["汇总", "明细"]:
                raise Failure("Unexpected workbook sheets")
            detail, summary = book["明细"], book["汇总"]
            end = len(ordered) + 2
            check(detail.max_row == end and detail.max_column == 9, "row_completeness", "Unexpected detail extent")
            check([detail.cell(1, column).value for column in range(1, 10)] == HEADERS, "row_completeness", "Header mismatch")
            check({str(merged) for merged in detail.merged_cells.ranges} == {f"A{end}:D{end}"},
                  "row_completeness", "Unexpected merged cells")
            check(len(manifest.rows) == len(ordered), "row_completeness", "Manifest row count mismatch")
            numbers, dates, attachments, calculated = [], [], set(), {}
            for index, item in enumerate(ordered, 2):
                invoice, rule = item.invoice, policy.categories[item.category]
                values = [detail.cell(index, column).value for column in range(1, 10)]
                amount = actual_cents(detail.cell(index, 5))
                check(amount == invoice.amount_cents.value, "row_completeness", "Amount differs from snapshot")
                detail_text = item.expense_detail
                if item.replaces_invoice_no:
                    detail_text += f"；重开票，原票 {item.replaces_invoice_no} 已作废"
                expected = [index - 1, neutral(rule.btype), neutral(rule.summary), neutral(detail_text), None,
                            item.service_date, invoice.invoice_no.value, neutral(invoice.seller_name.value)]
                check(all(values[n] == expected[n] for n in range(8) if n != 4), "row_completeness", "Row differs from snapshot")
                check(all(detail.cell(index, n).data_type == "s" for n in (2, 3, 4, 6, 7, 8, 9)),
                      "row_completeness", "Untrusted text is not a text cell")
                check(detail.cell(index, 7).number_format == "@", "row_completeness", "Invoice number format changed")
                number, filename = values[6], values[8]
                check(isinstance(number, str) and re.fullmatch(r"[0-9]{20}", number) is not None,
                      "unique_invoices", "Invalid invoice cell")
                numbers.append(number)
                dates.append(values[5])
                if not isinstance(filename, str) or "/" in filename or "\\" in filename or filename in RECORDS or filename.startswith("."):
                    raise Failure("Unsafe attachment name")
                attachments.add(filename)
                attachment = inside(directory / filename, store.root, must_exist=False)
                check(attachment.is_file() and hashlib.sha256(attachment.read_bytes()).hexdigest() == invoice.source_file_id,
                      "table_disk", "Attachment missing or changed")
                match = re.fullmatch(r"([0-9]{20})_(.+)_([0-9]+(?:\.[0-9]{1,2})?)_([^_]+)\.pdf", filename)
                check(match is not None and match[1] == number == invoice.invoice_no.value and match[2] == item.short_name
                      and decimal_cents(match[3]) == amount and match[4] == item.category
                      and match[3] == amount_text(amount), "filename_fields", "Filename differs from cells")
                calculated[rule.summary] = sum_cents([calculated.get(rule.summary, 0), amount])
                row = manifest.rows[index - 2] if index - 2 < len(manifest.rows) else None
                if row is None:
                    continue
                fields = row.model_dump()
                check(row.business_hash == digest({key: value for key, value in fields.items() if key != "business_hash"}),
                      "row_completeness", "Manifest row hash changed")
                check((row.item_id, row.invoice_no, row.amount_cents, row.service_date, row.relative_name, row.seller_name,
                       row.expense_detail, row.btype, row.summary, row.attachment_hash) ==
                      (item.id, number, amount, values[5], filename, invoice.seller_name.value, detail_text, rule.btype,
                       rule.summary, invoice.source_file_id), "row_completeness", "Manifest business fields differ")
            check(dates == sorted(dates), "date_order", "Dates are not sorted")
            check(len(numbers) == len(set(numbers)), "unique_invoices", "Duplicate invoice number in batch")
            check(not set(numbers) & history_numbers, "unique_invoices", "Invoice number already in history")
            disk = {entry.name for entry in directory.iterdir()}
            check(disk == attachments | {manifest.ledger_name} | RECORDS, "table_disk", "Disk inventory differs")
            check(len(manifest.files) == len(attachments) + 1 and
                  {entry.relative_name for entry in manifest.files} == attachments | {manifest.ledger_name},
                  "table_disk", "Manifest inventory differs")
            for entry in manifest.files:
                file = inside(directory / entry.relative_name, store.root, must_exist=False)
                check(file.is_file() and file.stat().st_size == entry.bytes
                      and hashlib.sha256(file.read_bytes()).hexdigest() == entry.sha256, "table_disk", "Artifact hash differs")
            total = sum_cents(calculated.values())
            check(summary.max_row <= max(4, len(calculated) + 2) and summary.max_column == 4,
                  "category_totals", "Summary extent differs")
            check(summary["D2"].value == neutral(snapshot.applicant), "row_completeness", "Applicant differs")
            check([summary.cell(1, 1).value, summary.cell(1, 2).value] == ["报销类型", "金额"],
                  "category_totals", "Summary header differs")
            for index, (category, amount) in enumerate(sorted(calculated.items()), 2):
                check(summary.cell(index, 1).value == neutral(category) and actual_cents(summary.cell(index, 2)) == amount,
                      "category_totals", "Category total differs")
            check(detail.cell(end, 1).value == "合计" and actual_cents(detail.cell(end, 5)) == total,
                  "category_totals", "Detail total differs")
            check(summary.cell(len(calculated) + 2, 1).value == "合计"
                  and actual_cents(summary.cell(len(calculated) + 2, 2)) == total,
                  "category_totals", "Summary total differs")
            check(manifest.total_cents == total == expected_total and manifest.category_totals == calculated,
                  "category_totals", "Manifest totals differ")
            check(not any(cell.data_type == "f" for sheet in book for row in sheet.iter_rows() for cell in row),
                  "row_completeness", "Unexpected formula")
            # Display formats, hidden rows and stray text are invisible to value checks: compare the whole file.
            rows = expected_rows(snapshot, policy)
            totals = {}
            for row in rows:
                totals[row["summary"]] = sum_cents([totals.get(row["summary"], 0), row["amount_cents"]])
            check(path.read_bytes() == workbook(rows, totals, snapshot.applicant), "row_completeness",
                  "Workbook differs from the snapshot rendering")
        finally:
            book.close()
    except Failure as failure:
        failures["row_completeness"].append(str(failure))
    except (CoreError, ValueError, TypeError, KeyError, IndexError, OSError, zipfile.BadZipFile):
        # pydantic ValidationError is a ValueError; any unreadable artifact fails closed.
        failures["row_completeness"].append("Unreadable or inconsistent artifacts")
    checks = [{"name": name, "passed": not failures[name], "issues": failures[name]} for name in CHECKS]
    return {"checks": checks, "passed": all(item["passed"] for item in checks),
            "issues": [issue for values in failures.values() for issue in values]}
