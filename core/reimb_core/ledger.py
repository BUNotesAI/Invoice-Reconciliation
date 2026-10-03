"""Pure rendering of the handover workbook. Package writes it; verify rebuilds it to compare bytes."""
import io
import re
import zipfile
from datetime import datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .values import sum_cents

HEADERS = ["序号", "报销类型", "项目", "费用明细", "金额", "日期", "发票号", "公司全称", "备注"]
SUMMARY_HEADERS = ["报销类型", "金额"]
FIXED_TIME = datetime(2026, 1, 1)


def spreadsheet_text(value):
    """Text written by users or agents is neutralised so no spreadsheet treats it as a formula."""
    value = str(value)
    return "'" + value if value.startswith(("=", "+", "-", "@")) else value


def normalize_xlsx(data):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            content = source.read(name)
            if name == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]+", rb"\g<1>2026-01-01T00:00:00Z", content)
            info = zipfile.ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, content)
    return output.getvalue()


def workbook(rows, totals, applicant):
    book = Workbook()
    summary = book.active
    summary.title = "汇总"
    summary.append(SUMMARY_HEADERS)
    for category, value in sorted(totals.items()):
        summary.append([spreadsheet_text(category), Decimal(value) / 100])
    summary.append(["合计", Decimal(sum_cents(totals.values())) / 100])
    summary["D1"] = "申请人"
    summary["D2"] = spreadsheet_text(applicant)
    summary["D4"] = "金额按确认快照生成；修改后须重新终审。"
    detail = book.create_sheet("明细")
    detail.append(HEADERS)
    for index, row in enumerate(rows, 1):
        values = [index, row["btype"], row["summary"], row["expense_detail"], Decimal(row["amount_cents"]) / 100,
                  row["service_date"], row["invoice_no"], row["seller_name"], row["relative_name"]]
        for column, value in enumerate(values, 1):
            cell = detail.cell(index + 1, column, spreadsheet_text(value) if isinstance(value, str) else value)
            if isinstance(value, str):
                cell.data_type = "s"
            if column == 7:
                cell.number_format = "@"
    end = len(rows) + 2
    detail.cell(end, 1, "合计")
    detail.merge_cells(start_row=end, start_column=1, end_row=end, end_column=4)
    detail.cell(end, 5, Decimal(sum_cents(row["amount_cents"] for row in rows)) / 100)
    border = Border(*(Side(style="thin", color="D3DED9"),) * 4)
    for sheet in book:
        sheet.freeze_panes = "A2"
        for cells in sheet:
            for cell in cells:
                cell.font = Font(name="Arial", size=11, bold=cell.row == 1)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = border
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="D9EAE2")
        sheet.row_dimensions[1].height = 26
    for column, width in {"A": 8, "B": 16, "C": 22, "D": 55, "E": 15, "F": 16, "G": 26, "H": 36, "I": 65}.items():
        detail.column_dimensions[column].width = width
    summary.column_dimensions["A"].width = 24
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["D"].width = 58
    for row in range(2, end + 1):
        detail.cell(row, 5).number_format = "#,##0.00"
    for row in range(2, len(totals) + 3):
        summary.cell(row, 2).number_format = "#,##0.00"
    book.properties.created = FIXED_TIME
    book.properties.modified = FIXED_TIME
    book.properties.creator = "Invoice Reconciliation"
    stream = io.BytesIO()
    book.save(stream)
    return normalize_xlsx(stream.getvalue())
