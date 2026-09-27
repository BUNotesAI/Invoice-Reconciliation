"""Content-addressed input objects and bounded, atomic local writes."""
import hashlib
import io
import os
import re
import signal
import stat
import tempfile
import zipfile
from contextlib import contextmanager
from pathlib import Path

import pypdf.filters
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader

from .errors import CoreError, require
from .models import SourceFile
from .values import canonical, strict_json

MAX_FILE = 20 * 1024 * 1024
MAX_FILES = 100
MAX_PAGES = 20
MAX_PIXELS = 40_000_000
MAX_STREAM = 16 * 1024 * 1024
MAX_XLSX_UNPACKED = 32 * 1024 * 1024
PARSE_SECONDS = 20
WECHAT_HEADER = ("交易时间", "交易类型", "交易对方", "商品", "收/支", "金额(元)", "支付方式", "当前状态", "交易单号", "商户单号", "备注")
DIDI_HEADER = ("序号", "车型", "上车时间", "城市", "起点", "终点", "里程", "金额")

# Bound decompression inside pypdf before any content is materialised.
pypdf.filters.ZLIB_MAX_OUTPUT_LENGTH = MAX_STREAM
pypdf.filters.MAX_DECLARED_STREAM_LENGTH = MAX_STREAM
Image.MAX_IMAGE_PIXELS = MAX_PIXELS


def inside(path, root, must_exist=True):
    path, root = Path(path), Path(root).resolve()
    require(path.is_absolute(), "INVALID_PATH", "Absolute path required")
    require(".." not in path.parts, "INVALID_PATH", "Parent path segments are not allowed")
    require(path == root or root in path.parents, "INVALID_PATH", "Path is outside runtime root")
    require(not root.is_symlink(), "INVALID_PATH", "Symbolic runtime root")
    cursor = path
    while cursor != root:
        require(not cursor.is_symlink(), "INVALID_PATH", "Symbolic links are not allowed")
        cursor = cursor.parent
    resolved = path.resolve()
    require(resolved == root or root in resolved.parents, "INVALID_PATH", "Path escapes runtime root")
    if must_exist:
        require(path.exists(), "INVALID_PATH", "Object does not exist")
    return path


def atomic_write(path, data):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    require(not path.is_symlink(), "INVALID_PATH", "Symbolic output object")
    fd, temporary = tempfile.mkstemp(prefix=".write-", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as handle:
            handle.write(data)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
        directory = os.open(path.parent, os.O_RDONLY)
        try:
            os.fsync(directory)
        finally:
            os.close(directory)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


@contextmanager
def parse_deadline():
    """Wall-clock bound for one file. It limits time, not memory; the byte limits above bound memory."""
    def timeout(_signum, _frame):
        raise CoreError("PARSE_LIMIT", "File parsing exceeded time limit")
    previous = signal.signal(signal.SIGALRM, timeout)
    signal.setitimer(signal.ITIMER_REAL, PARSE_SECONDS)
    try:
        yield
    finally:
        signal.setitimer(signal.ITIMER_REAL, 0)
        signal.signal(signal.SIGALRM, previous)


MAX_FORM_DEPTH = 8


def _pixels(width, height):
    require(0 <= int(width) * int(height) <= MAX_PIXELS, "PARSE_LIMIT", "Embedded image pixel limit exceeded")


def _inline_images(stream_owner, reader):
    """Declared size of every inline image (BI ... ID ... EI) in one content stream."""
    from pypdf.generic import ContentStream
    content = ContentStream(stream_owner, reader)
    for operands, operator in content.operations:
        if operator == b"INLINE IMAGE":
            settings = operands.get("settings", {})
            _pixels(settings.get("/W", settings.get("/Width", 0)), settings.get("/H", settings.get("/Height", 0)))


def _resource_images(resources, reader, seen, depth):
    """Image XObjects at any depth of Form XObject nesting, and inline images inside those forms."""
    require(depth <= MAX_FORM_DEPTH, "PARSE_LIMIT", "PDF form nesting limit exceeded")
    xobjects = resources.get_object().get("/XObject") if resources is not None else None
    if xobjects is None:
        return
    for reference in xobjects.get_object().values():
        key = getattr(reference, "idnum", None)
        if key is not None:
            if key in seen:
                continue
            seen.add(key)
        xobject = reference.get_object()
        subtype = xobject.get("/Subtype")
        if subtype == "/Image":
            _pixels(xobject.get("/Width", 0), xobject.get("/Height", 0))
        elif subtype == "/Form":
            _inline_images(xobject, reader)
            _resource_images(xobject.get("/Resources"), reader, seen, depth + 1)


def _image_pixels(page, reader):
    _resource_images(page.get("/Resources"), reader, set(), 0)
    if page.get("/Contents") is not None:
        _inline_images(page.get_contents(), reader)


def pdf_pages(data):
    """Text of each page, with page count, stream size and embedded image size bounded."""
    try:
        with parse_deadline():
            reader = PdfReader(io.BytesIO(data), strict=True)
            require(not reader.is_encrypted, "UNSUPPORTED_FILE", "Encrypted PDF is unsupported")
            require(0 < len(reader.pages) <= MAX_PAGES, "PARSE_LIMIT", "PDF page limit exceeded")
            texts, total = [], 0
            for page in reader.pages:
                contents = page.get_contents()
                require(contents is None or len(contents.get_data()) <= MAX_STREAM,
                        "PARSE_LIMIT", "PDF content limit exceeded")
                _image_pixels(page, reader)
                text = page.extract_text() or ""
                total += len(text)
                require(total <= 1024 * 1024, "PARSE_LIMIT", "PDF text limit exceeded")
                texts.append(text)
            return texts
    except CoreError:
        raise
    except pypdf.errors.LimitReachedError:
        raise CoreError("PARSE_LIMIT", "PDF stream limit exceeded") from None
    except Exception:
        raise CoreError("UNSUPPORTED_FILE", "Unreadable PDF") from None


def _has_row(text, header):
    # Column titles may carry a unit suffix such as 里程(公里).
    pattern = r"(?:[(（][^)）\n]*[)）])?[ \t]*".join(re.escape(word) for word in header)
    return re.search(pattern, text) is not None


def _pdf_kind(data):
    text = "\n".join(pdf_pages(data))
    if not text.strip():
        return "image_invoice_pdf"
    if "发票号码" in text and "价税合计" in text:
        return "invoice_pdf"
    if "行程单" in text and _has_row(text, DIDI_HEADER):
        return "didi_trip_pdf"
    return "unsupported"


def xlsx_rows(data, max_rows=5000):
    """Values of the first worksheet, after bounding the archive before openpyxl expands it."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            require(len(entries) <= 200 and sum(entry.file_size for entry in entries) <= MAX_XLSX_UNPACKED,
                    "PARSE_LIMIT", "Workbook archive limit exceeded")
        from openpyxl import load_workbook
        with parse_deadline():
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
            try:
                sheet = book.worksheets[0]
                rows = []
                for row in sheet.iter_rows(values_only=True):
                    rows.append(tuple(row))
                    require(len(rows) <= max_rows, "PARSE_LIMIT", "Workbook row limit exceeded")
                return rows
            finally:
                book.close()
    except CoreError:
        raise
    except Exception:
        raise CoreError("UNSUPPORTED_FILE", "Unreadable workbook") from None


def wechat_header_row(rows):
    for index, row in enumerate(rows[:60]):
        cells = tuple("" if value is None else str(value).strip() for value in row)
        if cells[:len(WECHAT_HEADER)] == WECHAT_HEADER:
            return index
    return None


def detect(data):
    if data.startswith(b"%PDF-"):
        return _pdf_kind(data)
    if data.startswith((b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff")):
        try:
            with parse_deadline(), Image.open(io.BytesIO(data)) as image:
                require(image.width * image.height <= MAX_PIXELS, "PARSE_LIMIT", "Image pixel limit exceeded")
                image.verify()
            return "image"
        except (Image.DecompressionBombError, Image.DecompressionBombWarning):
            raise CoreError("PARSE_LIMIT", "Image pixel limit exceeded") from None
        except (UnidentifiedImageError, OSError, SyntaxError):
            raise CoreError("UNSUPPORTED_FILE", "Unreadable image") from None
    if data.startswith(b"PK\x03\x04"):
        try:
            rows = xlsx_rows(data)
        except CoreError as error:
            if error.code == "PARSE_LIMIT":
                raise
            return "unsupported"
        return "wechat_bill" if wechat_header_row(rows) is not None else "unsupported"
    return "unsupported"


def read_upload(path):
    """Read through one descriptor opened without following links; a hard link could import a file from outside."""
    try:
        fd = os.open(path, os.O_RDONLY | os.O_NOFOLLOW)
    except OSError:
        raise CoreError("INVALID_PATH", "Input must be a regular file") from None
    with os.fdopen(fd, "rb") as handle:
        status = os.fstat(handle.fileno())
        require(stat.S_ISREG(status.st_mode) and status.st_nlink == 1, "INVALID_PATH", "Input must be a single-link regular file")
        require(status.st_size <= MAX_FILE, "INPUT_TOO_LARGE", "File size limit exceeded")
        data = handle.read(MAX_FILE + 1)
    require(len(data) <= MAX_FILE, "INPUT_TOO_LARGE", "File size limit exceeded")
    return data


class Store:
    def __init__(self, batch_dir, runtime_root):
        self.root = Path(runtime_root).resolve()
        self.batch = inside(batch_dir, self.root, must_exist=False)
        self.batch.mkdir(parents=True, exist_ok=True)
        self.objects = inside(self.batch / "objects", self.root, must_exist=False)
        self.objects.mkdir(exist_ok=True)

    def object_path(self, key):
        require(isinstance(key, str) and re.fullmatch(r"[0-9a-f]{64}", key), "INVALID_PATH", "Invalid object identifier")
        return inside(self.objects / key, self.root)

    def source(self, key):
        path = self.object_path(key)
        metadata = inside(self.objects / (key + ".json"), self.root)
        record = SourceFile.model_validate(strict_json(metadata.read_bytes()))
        require(hashlib.sha256(path.read_bytes()).hexdigest() == record.sha256 == key,
                "SNAPSHOT_MISMATCH", "Source content changed", 3)
        return record, path

    def ingest(self, source_path, original_name):
        require(isinstance(source_path, str) and isinstance(original_name, str) and len(original_name) <= 255
                and not re.search(r"[\x00-\x1f\x7f]", original_name), message="Invalid ingest input")
        path = inside(source_path, self.root)
        data = read_upload(path)
        key = hashlib.sha256(data).hexdigest()
        target = inside(self.objects / key, self.root, must_exist=False)
        metadata = inside(self.objects / (key + ".json"), self.root, must_exist=False)
        if metadata.exists():
            record, _ = self.source(key)
            return {"source_file": record.model_dump(), "duplicate": True}
        # Metadata is the commit marker; an object without it is a crash leftover and is rewritten.
        registered = sum(1 for _ in self.objects.glob("*.json"))
        require(registered < MAX_FILES, "FILE_LIMIT", "Batch file limit exceeded")
        kind = detect(data)
        record = SourceFile(id=key, sha256=key, byte_length=len(data), detected_type=kind,
                            original_name=original_name, storage_object_id=key)
        atomic_write(target, data)
        atomic_write(metadata, canonical(record.model_dump()))
        return {"source_file": record.model_dump(), "duplicate": False}
